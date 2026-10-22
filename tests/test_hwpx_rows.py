"""세부내역 행 늘리기(clone)·묶기(group), reader 왕복, 예시 문서 읽기."""
import re
import xml.etree.ElementTree as ET
import zipfile
from datetime import date, datetime

import pytest

from core.ledger import parse_date
from core.models import LineItem, ReportDoc, RequestDoc
from core.paths import TEMPLATES_DIR
from core.validate import validate_report
from hwpx.fill import REPORT_TEMPLATE, REQUEST_TEMPLATE, item_rows, render
from hwpx.layout import REPORT, resolve
from hwpx.reader import UnknownFormError, read_document
from hwpx.table import expand_item_rows
from hwpx.text import cells_from_xml, extract_text, read_cells
from tests.conftest import TODAY

EXAMPLES = TEMPLATES_DIR / "examples"
ACCOUNT = "초등부 교사회"


def items(n, account=ACCOUNT):
    return [LineItem(계정과목=account, 내용=f"{i + 1}번 항목 식대", 단가=1_000 * (i + 1), 수량=(i % 3) + 1,
                     비고="영수증" if i % 2 else "") for i in range(n)]


def full_fields(month=8):
    return dict(월=month, 사업ID="P11", 행사명=f"{month}월 교사회 운영비", 작성일=date(2026, month, 30),
                주관단체=ACCOUNT, 협조단체="청소년분과", 기안자="정연재 레오", 기안단체장="정지연 플로라",
                기안자연락처="010-0000-0001", 기안단체장연락처="010-0000-0002",
                은행명="테스트은행", 계좌번호="000-00-0000", 예금주="홍길동",
                목적="교사회 운영", 시작=datetime(2026, month, 1, 10), 종료=datetime(2026, month, 31, 18, 30),
                박=1, 일=2, 장소="개포동성당", 참가대상="초등부 교사회", 인원_총=9, 인원_남=5, 인원_여=4)


def built_request(ledger, n):
    doc = RequestDoc(**full_fields(), 행사일정="첫 줄 설명\n\n- 둘째 줄 & <특수문자>", 세부내역=items(n))
    return ledger.build_request(doc, today=TODAY).docs[0]


def built_report(ledger, n):
    doc = ReportDoc(**full_fields(), 경과보고="다음과 같이 사용하였습니다.\n- 식대", 건의사항="다음 달 인원 증가",
                    세부내역=items(n))
    return ledger.build_report(doc, ledger.row(53), today=TODAY).docs[0]


IGNORED = {"연결문서번호", "양식", "행사유형", "계좌_override"}


def comparable(doc):
    return doc.model_dump(exclude=IGNORED)


# ---------------------------------------------------------------- clone

@pytest.mark.parametrize("n", [1, 7, 12])
@pytest.mark.parametrize("kind", ["request", "report"])
def test_round_trip(ledger, tmp_path, n, kind):
    doc = built_request(ledger, n) if kind == "request" else built_report(ledger, n)
    path = render(doc, tmp_path / f"{kind}{n}.hwpx")
    res = read_document(path, 사업ID="P11")
    assert res.failures == []
    assert res.item_rows == max(7, n)
    assert comparable(res.doc) == comparable(doc)


def test_twelve_items_structure(ledger, tmp_path):
    doc = built_report(ledger, 12)
    path = render(doc, tmp_path / "r.hwpx")
    with zipfile.ZipFile(path) as z, zipfile.ZipFile(REPORT_TEMPLATE) as t:
        xml, base = z.read("Contents/section0.xml").decode(), t.read("Contents/section0.xml").decode()
        for i in z.infolist():
            if i.filename.endswith((".xml", ".hpf", ".rdf")):
                ET.fromstring(z.read(i))
        assert z.read("Preview/PrvText.txt").decode() == extract_text(path)
    assert re.search(r'rowCnt="(\d+)"', xml).group(1) == "32"
    assert xml.count("<hp:tc ") == base.count("<hp:tc ") + 5 * 6
    height = lambda x: int(re.search(r'<hp:sz width="\d+" widthRelTo="ABSOLUTE" height="(\d+)"', x).group(1))
    assert height(xml) == height(base) + 5 * 1734
    assert 'pageBreak="CELL"' in re.search(r"<hp:tbl [^>]*>", xml).group()
    cells = read_cells(path)
    assert cells[(27, 1)] == ["12번 항목 식대"]
    assert cells[(28, 1)][0].replace(" ", "") == "합계(B)" and cells[(28, 13)] == [f"{doc.합계B:,}원"]
    assert cells[(29, 13)] == [f"{doc.잔액:,}원"]
    assert "12번 항목 식대" in extract_text(path)


def test_expand_adjusts_merged_span():
    """세부내역 위에서 내려오는 병합 칸(예시 문서의 계정과목 칸)은 rowSpan·높이가 함께 늘어난다."""
    xml = zipfile.ZipFile(REPORT_TEMPLATE).read("Contents/section0.xml").decode()
    xml = re.sub(r'<hp:tc\b(?:(?!</hp:tc>).)*?colAddr="0" rowAddr="(1[7-9]|2[0-2])"(?:(?!</hp:tc>).)*?</hp:tc>',
                 "", xml, flags=re.S)  # 17~22행의 계정과목 칸 삭제
    xml = re.sub(r'(colAddr="0" rowAddr="16"/><hp:cellSpan colSpan="1" rowSpan=")1("/><hp:cellSz width="\d+" height=")(\d+)',
                 lambda m: f"{m.group(1)}7{m.group(2)}{int(m.group(3)) * 7}", xml)
    out = expand_item_rows(xml, 3)
    tc = re.search(r'<hp:tc\b(?:(?!</hp:tc>).)*?colAddr="0" rowAddr="16"(?:(?!</hp:tc>).)*?</hp:tc>', out, re.S).group()
    assert 'rowSpan="10"' in tc
    assert re.search(r'height="(\d+)"', tc.split("<hp:cellSz ")[1]).group(1) == str(1734 * 7 + 1734 * 3)
    res = resolve(cells_from_xml(out), REPORT)
    assert (res.first_item_row, res.total_row) == (16, 26)


# ---------------------------------------------------------------- group

def test_group_mode_fits_seven_rows(ledger, tmp_path):
    doc = built_report(ledger, 12)
    rows = item_rows(doc, "group")
    assert len(rows) == 7
    assert rows[-1]["내용"] == "7번 항목 식대 외 5건" and rows[-1]["단가"] == rows[-1]["수량"] == ""
    assert sum(int(r["금액"].replace(",", "")) for r in rows) == doc.합계B
    path = render(doc, tmp_path / "g.hwpx", row_mode="group")
    res = read_document(path)
    assert res.item_rows == 7 and len(res.doc.세부내역) == 7
    assert sum(i.금액 for i in res.doc.세부내역) == doc.합계B


def test_group_mode_merges_largest_account_first(ledger):
    its = items(6, "가") + items(3, "나")  # 9건 → 2건 줄여야 함: '가'에서 묶음
    doc = ReportDoc(월=8, 사업ID="P11", 행사명="x", 세부내역=its)
    rows = item_rows(doc, "group")
    assert [r["계정과목"] for r in rows] == ["가"] * 4 + ["나"] * 3
    assert [r["내용"] for r in rows] == ["1번 항목 식대", "2번 항목 식대", "3번 항목 식대", "4번 항목 식대 외 2건",
                                        "1번 항목 식대", "2번 항목 식대", "3번 항목 식대"]


def test_group_mode_impossible(ledger):
    its = [LineItem(계정과목=f"계정{i}", 내용=f"항목{i}", 단가=1_000) for i in range(8)]
    doc = ReportDoc(월=8, 사업ID="P11", 행사명="8월 교사회 운영비", 세부내역=its)
    with pytest.raises(ValueError):
        item_rows(doc, "group")
    ledger.settings.hwpx_row_mode = "group"
    assert "too_many_items" in {f.코드 for f in validate_report(doc, ledger, today=TODAY)}
    ledger.settings.hwpx_row_mode = "clone"
    assert "too_many_items" not in {f.코드 for f in validate_report(doc, ledger, today=TODAY)}


# ---------------------------------------------------------------- reader

def test_reader_blank_template_and_other_documents(tmp_path):
    blank = read_document(REQUEST_TEMPLATE)  # 빈 양식: 서식은 판별되고, 작성일·문서번호 없음이 실패로 남는다
    assert blank.kind == "예산신청서" and blank.doc.세부내역 == []
    assert {f.칸 for f in blank.failures} == {"작성일", "문서번호"}
    path = tmp_path / "other.hwpx"
    with zipfile.ZipFile(REQUEST_TEMPLATE) as src, zipfile.ZipFile(path, "w") as dst:
        for info in src.infolist():
            data = src.read(info)
            if info.filename == "Contents/section0.xml":
                data = data.decode().replace("<hp:t>예 산 신 청 서</hp:t>", "<hp:t>회 의 록</hp:t>").encode()
            dst.writestr(info, data, compress_type=info.compress_type)
    with pytest.raises(UnknownFormError):
        read_document(path)


def test_reader_reports_failures(ledger, tmp_path):
    doc = built_request(ledger, 2)
    path = render(doc, tmp_path / "a.hwpx")
    with zipfile.ZipFile(path) as src, zipfile.ZipFile(tmp_path / "b.hwpx", "w") as dst:
        for info in src.infolist():
            data = src.read(info)
            if info.filename == "Contents/section0.xml":
                data = data.decode().replace(" 2026년 8월 30일", "이천이십육년").replace(
                    f"{doc.합계:,}원", "약 3만원").encode()
            dst.writestr(info, data, compress_type=info.compress_type)
    res = read_document(tmp_path / "b.hwpx")
    assert {f.칸 for f in res.failures} >= {"작성일", "합계"}


@pytest.mark.parametrize("name", ["request_childrensday", "report_childrensday",
                                  "request_camp_aug", "report_camp_aug"])
def test_examples_match_ledger(ledger, name):
    path = EXAMPLES / f"{name}.hwpx"
    if not path.exists():
        pytest.skip("작성 예시 없음(커밋 제외 파일)")
    res = read_document(path)
    d = res.doc
    assert res.failures == []
    row = next(r for r in ledger.rows if r.문서번호 == d.문서번호)
    assert parse_date(row.작성일) == d.작성일 and row.행사명 == d.행사명
    assert (row.자체예산, row.본당보조) == (d.자체예산, d.본당보조)
    if res.kind == "예산신청서":
        assert (row.예산배정액, row.금번신청액, row.신청누계액, row.잔여액, row.합계) == \
            (d.예산배정액, d.금번신청액, d.신청누계액, d.잔여액, d.합계)
    else:
        assert (row.합계, row.잔액) == (d.합계B, d.잔액)
        assert d.예산A == ledger.find_ref(row.연결문서번호).금번신청액 + d.자체예산


# ---------------------------------------------------------------- 영수증 docx

def test_receipts_docx(tmp_path):
    from docx import Document
    from PIL import Image

    from core.models import Receipt
    from hwpx.receipts import build_receipts_docx

    rotated = Image.new("RGB", (400, 200), "white")
    rotated.paste(Image.new("RGB", (300, 100), "black"), (50, 50))
    exif = Image.Exif()
    exif[0x0112] = 6  # 90도 회전해서 볼 사진
    rotated.save(tmp_path / "rot.jpg", exif=exif)
    Image.new("RGB", (200, 300), "gray").save(tmp_path / "scan.pdf")
    receipts = [Receipt(순번=2, 이미지="scan.pdf", 날짜=date(2026, 8, 16), 상호="가게B", 금액=2_000, 용도="간식"),
                Receipt(순번=1, 이미지="rot.jpg", 날짜=date(2026, 8, 15), 상호="가게A", 금액=10_000, 용도="식대")]
    path = build_receipts_docx("(집)2026-099", "테스트 행사", receipts, tmp_path / "r.docx", base_dir=tmp_path)

    doc = Document(path)
    assert doc.paragraphs[0].text.strip() == "(집)2026-099  테스트 행사"
    assert "합계 12,000원" in doc.paragraphs[1].text
    assert [[c.text for c in r.cells] for r in doc.tables[0].rows] == [
        ["날짜", "2026-08-15"], ["상호", "가게A"], ["금액", "10,000원"], ["용도", "식대"]]  # 순번순
    with zipfile.ZipFile(path) as z:
        media = sorted(n for n in z.namelist() if n.startswith("word/media/"))
        first = Image.open(z.open(media[0]))
    assert len(media) == 2
    assert first.height > first.width  # EXIF 회전 반영(400x200 → 세로)
    assert first.size[1] < 400 and first.size[0] < 200  # 흰 여백 잘라냄
