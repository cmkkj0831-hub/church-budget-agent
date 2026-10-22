from openpyxl import load_workbook

from core.export import export_xlsx


def test_export_xlsx(ledger, book, tmp_path):
    path = export_xlsx(ledger, book, tmp_path / "out.xlsx")
    wb = load_workbook(path)
    ws = wb["2026 신청·집행 내역"]
    header = [c.value for c in ws[1]]
    assert header[:3] == ["순번", "월", "구분"] and "비고" in header
    assert ws.max_row >= 54
    assert ws.freeze_panes == "E2" and ws.auto_filter.ref.startswith("A1:")
    note_col = header.index("비고") + 1
    assert ws.cell(2, 3).value == "예산신청서" and ws.cell(2, 1).fill.fgColor.rgb.endswith("FFFFFF")
    assert ws.cell(4, 3).value == "집행보고서" and ws.cell(4, 1).fill.fgColor.rgb.endswith("EAF1F8")
    assert ws.cell(2, note_col).value == "문서번호 미기재"
    assert ws.cell(2, note_col).font.color.rgb.endswith("C00000")
    assert ws.cell(20, note_col).value is None  # 순번 19: 이슈 해결됨
    assert header[-1] == "양식" and ws.cell(2, len(header)).value == "구"
    s = wb["누계"]
    assert [s.cell(r, 2).value for r in (2, 3, 4)] == [23_216_860, 23_694_010, -477_150]
    assert s.cell(4, 4).value == "열린 이슈로 모두 설명됨"
    rich = load_workbook(path, rich_text=True).active.cell(54, note_col).value
    colors = [b.font.color.rgb[-6:] for b in rich if hasattr(b, "font")]
    assert colors == ["C00000", "C65911"]  # 순번 53: 기재 오류 빨강, 예산 경고 주황
    wi = wb["이슈"]
    assert wi.max_row == len(book.issues) + 1
    assert [c.value for c in wi[1]][:2] == ["이슈ID", "대상순번"]
