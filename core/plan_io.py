"""사업계획서 읽기·검증·저장. 형식은 2026_사업계획서.xlsx와 같다('사업계획' 시트 + '개요·비고' 시트).

'사업계획' 시트: '단위' 행(천원/원), '사업ID' 머리 행, 사업 행(P01…), 열은 사업ID·시기·계속/신규·사업명·세부내역·
자체·본당·전년도 본당 집행액·1월~12월. 금액은 원 단위로 바꿔 data/{연도}/plan.csv에 쓴다.
"""
import re
import unicodedata
from pathlib import Path

import openpyxl

from core.ledger import write_csv

MONTHS = [f"{m}월" for m in range(1, 13)]
PLAN_COLUMNS = ["사업ID", "시기", "계속신규", "사업명", "세부내역", "자체예산", "본당예산", "전년도집행액", *MONTHS]
INFO_COLUMNS = ["항목", "내용"]


def _nfc(v) -> str:
    return unicodedata.normalize("NFC", str(v)) if v is not None else ""


def _num(v) -> float:
    if v in (None, ""):
        return 0
    if isinstance(v, (int, float)):
        return v
    s = re.sub(r"[^\d.\-]", "", str(v))
    return float(s) if s else 0


def parse_plan_xlsx(path) -> tuple[list[dict], list[dict]]:
    """xlsx → (사업 행 목록(원 단위), 개요·비고 목록)."""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["사업계획"] if "사업계획" in wb.sheetnames else wb.worksheets[0]
    unit, header, plan = 1000, None, []
    for row in ws.iter_rows(values_only=True):
        first = _nfc(row[0]).strip()
        if first == "단위":
            unit = 1 if "천" not in _nfc(row[1]) else 1000
        if first == "사업ID":
            header = [_nfc(h).strip() for h in row]
            continue
        if header is None or not re.fullmatch(r"P\d{2}", first):
            continue
        r = dict(zip(header, row))
        won = lambda v: int(round(_num(v) * unit))
        p = {"사업ID": first, "시기": _nfc(r.get("시기")), "계속신규": _nfc(r.get("계속/신규") or r.get("계속신규")),
             "사업명": _nfc(r.get("사업명")), "세부내역": _nfc(r.get("세부내역")),
             "자체예산": won(r.get("자체")), "본당예산": won(r.get("본당")),
             "전년도집행액": "" if r.get("전년도 본당 집행액") in (None, "") else won(r.get("전년도 본당 집행액"))}
        for m in MONTHS:
            p[m] = won(r.get(m))
        plan.append(p)
    info = []
    if "개요·비고" in wb.sheetnames:
        info = [{"항목": _nfc(a), "내용": _nfc(b)} for a, b, *_ in wb["개요·비고"].iter_rows(min_row=2, values_only=True)
                if a]
    return plan, info


def plan_problems(plan: list[dict], allocation: int | None = None) -> list[str]:
    """확정 전 검증: 사업ID 중복, 월별 배분 합 = 본당예산, (배정액이 있으면) 본당예산 합계 비교."""
    problems = []
    ids = [p["사업ID"] for p in plan]
    for dup in sorted({i for i in ids if ids.count(i) > 1}):
        problems.append(f"사업ID {dup} 중복")
    for p in plan:
        months = sum(int(p[m] or 0) for m in MONTHS)
        if months != int(p["본당예산"] or 0):
            problems.append(f"{p['사업ID']} {p['사업명']}: 월별 배분 합 {months:,}원 ≠ 본당예산 {int(p['본당예산'] or 0):,}원")
        if not p["사업명"]:
            problems.append(f"{p['사업ID']}: 사업명 없음")
    total = sum(int(p["본당예산"] or 0) for p in plan)
    if allocation is not None and total != allocation:
        problems.append(f"본당예산 합계 {total:,}원 ≠ 배정액 {allocation:,}원(배정액이 다르면 확인 후 진행)")
    return problems


def write_plan(year_dir, plan: list[dict], info: list[dict]) -> None:
    yd = Path(year_dir)
    yd.mkdir(parents=True, exist_ok=True)
    write_csv(yd / "plan.csv", PLAN_COLUMNS, plan)
    write_csv(yd / "plan_info.csv", INFO_COLUMNS, info)
