"""장부 + 이슈 → 회계내역 xlsx (원본 회계내역 서식)."""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from core.issues import (
    ISSUE_COLUMNS, IssueBook, cumulative_reconciliation, issue_color, issue_to_csv,
)
from core.ledger import Ledger

FONT = "맑은 고딕"
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
REQUEST_FILL = PatternFill("solid", fgColor="FFFFFF")
REPORT_FILL = PatternFill("solid", fgColor="EAF1F8")
RED = "C00000"
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
MONEY = "#,##0;-#,##0;0"

# (헤더, 장부 열 또는 '비고', 너비)
LEDGER_SHEET = [
    ("순번", "순번", 6), ("월", "월", 6), ("구분", "구분", 11), ("작성일", "작성일", 15),
    ("예산 배정액", "예산배정액", 13), ("금번 신청액", "금번신청액", 12),
    ("신청 누계액", "신청누계액", 13), ("잔여액", "잔여액", 13), ("자체 예산", "자체예산", 11),
    ("본당 보조", "본당보조", 12), ("합계", "합계", 12), ("잔액", "잔액", 10),
    ("문서번호", "문서번호", 14), ("행사명", "행사명", 30), ("비고", "비고", 60),
    ("파일명", "파일명", 34), ("사업ID", "사업ID", 7), ("연결문서번호", "연결문서번호", 14),
    ("상태", "상태", 10), ("양식", "양식", 6),
]
MONEY_COLS = {"예산배정액", "금번신청액", "신청누계액", "잔여액", "자체예산", "본당보조", "합계", "잔액"}
CENTER_COLS = {"순번", "월", "구분", "작성일", "문서번호", "사업ID", "연결문서번호", "상태", "양식"}
ISSUE_WIDTHS = {"설명": 50, "근거": 60, "해결내용": 40, "대상문서번호": 28}
STATE_FONT = {"미해결": RED, "확인중": "7F6000"}


def _header(ws, headers, widths):
    for c, (h, w) in enumerate(zip(headers, widths), 1):
        cell = ws.cell(1, c, h)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER
        ws.column_dimensions[get_column_letter(c)].width = w


def _remark(items: list[tuple[str, str]]):
    """비고: 열린 이슈 설명을 유형별 색으로(문서 오류 빨강, 예산 경고 주황)."""
    if not items:
        return None
    parts = []
    for n, (kind, desc) in enumerate(items):
        if n:
            parts.append(" / ")
        parts.append(TextBlock(InlineFont(rFont=FONT, sz=10, color=issue_color(kind)), desc))
    return CellRichText(parts)


def _cumulative_sheet(wb, ledger, book):
    rec = cumulative_reconciliation(ledger, book)
    s = rec.summary
    ws = wb.create_sheet("누계")
    rows = [
        ("구분", "신청누계액", "잔여액", "비고"),
        ("문서 기준", s.문서기준, s.문서기준_잔여, f"순번 {s.문서기준_순번} 신청서에 적힌 값"),
        ("재계산 기준", s.재계산, s.재계산_잔여, "신청액 합 − 반납액 합(실집행 기준, 취소 제외)"),
        ("차이(문서 − 재계산)", s.차이, -s.차이,
         "열린 이슈로 모두 설명됨" if rec.설명됨 else f"설명 안 된 순번: {rec.미설명}"),
        (),
        ("차이가 생긴 문서", "순번", "변화량", "이슈"),
    ]
    descs = {i.이슈ID: i.설명 for i in book.issues}
    for item in rec.항목:
        rows.append((item["문서번호"], item["순번"], item["차이"],
                     " / ".join(f"{i} {descs[i]}" for i in item["이슈"]) or "설명 없음"))
    for r, values in enumerate(rows, 1):
        for c, v in enumerate(values, 1):
            cell = ws.cell(r, c, v)
            cell.font = Font(name=FONT, size=10, bold=r in (1, 6))
            if isinstance(v, int) and not (r > 6 and c == 2):
                cell.number_format = MONEY
    for col, w in zip("ABCD", (20, 14, 14, 90)):
        ws.column_dimensions[col].width = w
    return ws


def _save(wb, path) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def export_xlsx(ledger: Ledger, book: IssueBook, path) -> Path:
    """회계내역: 신청·집행 내역(비고=열린 이슈) + 이슈 + 누계 시트."""
    wb = Workbook()
    write_ledger_sheets(wb, ledger, book, wb.active)
    return _save(wb, path)


def write_ledger_sheets(wb, ledger: Ledger, book: IssueBook, ws=None) -> None:
    remarks = book.ledger_remarks_detail()
    ws = ws or wb.create_sheet()
    ws.title = f"{ledger.year} 신청·집행 내역"
    _header(ws, [h for h, _, _ in LEDGER_SHEET], [w for _, _, w in LEDGER_SHEET])

    for i, r in enumerate(ledger.rows, 2):
        fill = REQUEST_FILL if r.신청서 else REPORT_FILL
        for c, (_, col, _) in enumerate(LEDGER_SHEET, 1):
            value = _remark(remarks.get(r.순번)) if col == "비고" else getattr(r, col)
            cell = ws.cell(i, c, value if value != "" else None)
            cell.fill = fill
            cell.border = BORDER
            cell.font = Font(name=FONT, size=10, color=RED if col == "비고" else None)
            if col in MONEY_COLS:
                cell.number_format = MONEY
            if col in CENTER_COLS:
                cell.alignment = Alignment(horizontal="center")
            if col == "비고":
                cell.alignment = Alignment(wrap_text=True, vertical="center")
    last = len(ledger.rows) + 1
    ws.auto_filter.ref = f"A1:{get_column_letter(len(LEDGER_SHEET))}{last}"
    ws.freeze_panes = "E2"
    note = ws.cell(last + 2, 1, "※ 흰색 행 = 예산신청서, 파란색 행 = 집행보고서. 비고는 열린 이슈의 설명이며 "
                                "해결되면 사라진다(빨강 = 문서 오류, 주황 = 예산 경고). 누계 두 기준은 '누계' 시트.")
    note.font = Font(name=FONT, size=9, color="595959")

    wi = wb.create_sheet("이슈")
    _header(wi, ISSUE_COLUMNS, [ISSUE_WIDTHS.get(c, 12) for c in ISSUE_COLUMNS])
    for i, issue in enumerate(book.issues, 2):
        d = issue_to_csv(issue)
        for c, col in enumerate(ISSUE_COLUMNS, 1):
            cell = wi.cell(i, c, d[col] or None)
            cell.border = BORDER
            color = (STATE_FONT.get(issue.상태) if col == "상태"
                     else issue_color(issue.유형) if col == "유형" else None)
            cell.font = Font(name=FONT, size=10, color=color)
            cell.alignment = Alignment(wrap_text=col in ISSUE_WIDTHS, vertical="center")
    wi.auto_filter.ref = f"A1:{get_column_letter(len(ISSUE_COLUMNS))}{len(book.issues) + 1}"
    wi.freeze_panes = "B2"
    _cumulative_sheet(wb, ledger, book)


def export_plan_xlsx(ledger: Ledger, plan_info: list[dict], path, today=None) -> Path:
    """사업계획서 + 집행 현황(사업별 실집행·집행률·남은 한도·연말 예상) + 개요·비고."""
    wb = Workbook()
    write_plan_sheets(wb, ledger, plan_info, today, wb.active)
    return _save(wb, path)


def write_plan_sheets(wb, ledger: Ledger, plan_info: list[dict], today=None, ws=None) -> None:
    ws = ws or wb.create_sheet()
    ws.title = "사업계획"
    months = [f"{m}월" for m in range(1, 13)]
    headers = ["사업ID", "시기", "계속신규", "사업명", "세부내역", "자체예산", "본당예산", "전년도집행액",
               "실집행", "집행률", "남은한도", "연말예상", *months]
    _header(ws, headers, [7, 6, 7, 22, 40, 11, 11, 11, 11, 8, 11, 11, *[9] * 12])
    for r, p in enumerate(ledger.plan.values(), 2):
        spent = ledger.business_spent(p.사업ID)
        rate = spent / p.본당예산 if p.본당예산 else 0
        row = [p.사업ID, p.시기, p.계속신규, p.사업명, p.세부내역, p.자체예산, p.본당예산, p.전년도집행액,
               spent, rate, ledger.business_limit(p.사업ID), ledger.year_end_projection(p.사업ID, today), *p.월별]
        for c, v in enumerate(row, 1):
            cell = ws.cell(r, c, v)
            cell.border = BORDER
            cell.font = Font(name=FONT, size=10)
            cell.number_format = "0.0%" if headers[c - 1] == "집행률" else MONEY
    total = len(ledger.plan) + 2
    ws.cell(total, 4, "합계").font = Font(name=FONT, bold=True)
    for c in (6, 7, 9, 11, 12, *range(13, 25)):
        col = get_column_letter(c)
        cell = ws.cell(total, c, f"=SUM({col}2:{col}{total - 1})")
        cell.number_format = MONEY
        cell.font = Font(name=FONT, bold=True)
    ws.freeze_panes = "E2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{total - 1}"

    wi = wb.create_sheet("개요·비고")
    _header(wi, ["항목", "내용"], [16, 100])
    for r, item in enumerate(plan_info, 2):
        wi.cell(r, 1, item["항목"]).font = Font(name=FONT, size=10)
        cell = wi.cell(r, 2, item["내용"])
        cell.font = Font(name=FONT, size=10)
        cell.alignment = Alignment(wrap_text=True, vertical="top")


def write_usage_sheet(wb, usage_rows: list[dict], pricing) -> None:
    from core.usage import summarize
    ws = wb.create_sheet("사용량")
    rows = summarize(usage_rows, pricing)
    headers = ["월", "기능", "환경", "호출", "입력토큰", "출력토큰", "소요초", "예상비용_원"]
    _header(ws, headers, [9, 16, 7, 7, 12, 12, 9, 12])
    for r, g in enumerate(rows, 2):
        for c, h in enumerate(headers, 1):
            ws.cell(r, c, g[h]).font = Font(name=FONT, size=10)


def export_annual_xlsx(ledger: Ledger, book: IssueBook, plan_info: list[dict], usage_rows: list[dict],
                       path, today=None) -> Path:
    """연간 결산: 장부·이슈·누계 + 사업별 집행 + 개요·비고 + 사용량."""
    wb = Workbook()
    write_ledger_sheets(wb, ledger, book, wb.active)
    write_plan_sheets(wb, ledger, plan_info, today)
    write_usage_sheet(wb, usage_rows, ledger.settings.사용량단가)
    return _save(wb, path)


if __name__ == "__main__":  # noqa: E305
    from core.paths import OUT_DIR
    out = export_xlsx(Ledger.load(), IssueBook.load(), OUT_DIR / "2026_초등부교사회_회계내역.xlsx")
    print(out)
