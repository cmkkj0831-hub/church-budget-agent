"""2026 엑셀(회계내역·사업계획서) → data/*.csv 1회 이관.

    python -m core.migrate          # data/ 파일이 이미 있으면 중단
    python -m core.migrate --force  # 덮어쓰기(이슈 처리 이력이 사라지므로 주의)

장부 값은 엑셀에 적힌 그대로 옮긴다. 비고 열은 옮기지 않고 이슈로 등록한다.
"""
import argparse
import json
import re
import sys

import openpyxl

from core.ledger import (
    HISTORY_COLUMNS, ISSUE_COLUMNS, LEDGER_COLUMNS, MULTI_SEP,
    parse_date, write_csv,
)
from core.paths import (
    DATA_DIR, HISTORY_CSV, ISSUES_CSV, LEDGER_CSV, PLAN_CSV, PLAN_INFO_CSV,
    SETTINGS_JSON, find_child, nfc, source_dir,
)

LEDGER_XLSX = "2026_초등부교사회_회계내역.xlsx"
PLAN_XLSX = "2026_사업계획서.xlsx"

# 엑셀 헤더 → 장부 열
LEDGER_HEADER_MAP = {
    "월": "월", "구분": "구분", "작성일": "작성일", "예산 배정액": "예산배정액",
    "금번 신청액": "금번신청액", "신청 누계액": "신청누계액", "잔여액": "잔여액",
    "자체 예산": "자체예산", "본당 보조": "본당보조", "합계": "합계", "잔액": "잔액",
    "문서번호": "문서번호", "행사명": "행사명", "비고": "비고", "파일명": "파일명",
}

# 행사명 → 사업ID. 위에서부터 첫 일치. confident=False면 확인 필요 목록에 올린다.
PROJECT_RULES = [
    (r"부서\s*운영비", "P13", True),
    (r"교사회\s*운영비", "P11", True),
    (r"교리\s*준비비", "P12", True),           # 부활 교리 준비비 포함
    (r"졸업", "P01", True),
    (r"동계", "P02", True),
    (r"성소", "P04", True),
    (r"어린이\s*날", "P05", True),
    (r"신앙\s*캠프|여름\s*캠프", "P07", True),  # 2월 선입금 포함
    (r"신앙학교\s*준비비", "P06", True),        # 6월 준비비: 8월 신앙학교 사전 준비(확정)
    (r"신앙학교", "P06", True),
    (r"하계", "P08", True),
]

# 이관 시 등록할 이슈. (대상순번 목록, 유형, 설명(=비고), 근거, 상태[, 관련이슈 목록])
INITIAL_ISSUES = [
    ([1], "기재 오류", "문서번호 미기재",
     "1_1월교사동계워크숍_신청서.hwp 문서번호란 공란", "미해결"),
    ([2], "기재 오류", "문서번호 미기재",
     "2_1월교사회운영비_신청서.hwp 문서번호란 공란", "미해결"),
    ([3], "기재 오류", "작성일 2025년으로 기재",
     "3_1월교사회운영비_보고서.hwp 작성일 '2025년 1월 31일', 일시란도 2025년 1월", "미해결"),
    ([4], "기재 오류", "작성일 2025년으로 기재",
     "4_1월교사동계워크숍_보고서.hwp 작성일 '2025년 2월 1일', 일시란 '2025년 2월 13일~14일'", "미해결"),
    ([9], "기재 오류", "작성일 2025년으로 기재",
     "5_2월교사회운영비_보고서.hwp 작성일 '2025년 2월 28일', 일시란도 2025년 2월", "미해결"),
    ([15], "기재 오류", "일시란 2025년으로 기재",
     "3_3월교사회운영비_보고서.hwp 작성일은 2026년 4월 5일이나 일시란 '2025년 3월 1일 ~ 3월 31일'", "미해결"),
    ([4, 5], "번호 중복", "문서번호 004가 1월 동계워크숍 보고서와 2월 운영비 신청서에 중복",
     "(집)2026-004 = 4_1월교사동계워크숍_보고서.hwp, (예)2026-004 = 1_2월교사회윤영비_신청서.hwp. "
     "신청서·보고서는 일련번호를 공유하므로 중복 불가", "미해결"),
    ([14], "계산 불일치", "합계란 700,000원 기재(세부내역 23,400원)",
     "2_3월교리준비비_신청서.hwp 세부내역 합 23,400원, 금번신청액·본당보조 23,400원, 합계란 700,000원", "미해결"),
    ([19], "기재 오류", "작성일 미기재",
     "회계내역 엑셀 작성일 공란. 단, 이관 중 원본 3_4월성소주일행사비_신청서.hwp 상단 날짜란에서 "
     "'2026 년 4 월 25 일'이 추출됨 → 장부 이관 누락 가능성, 원본 확인 필요", "미해결"),
    ([23], "기재 오류", "합계란 공란",
     "1_5월교사회운영비_신청서.hwpx 합계란 공란(금번신청액·본당보조 480,000원)", "미해결"),
    ([24], "계산 불일치", "세부내역 합 500,000원 ≠ 합계 490,000원",
     "2_어린이날행사비_신청서.hwpx 세부내역 합 500,000원, 합계·신청액 490,000원", "미해결"),
    ([25], "계산 불일치", "세부내역 합 490,070원 ≠ 합계 490,060원",
     "3_어린이날행사비_보고서.hwp 세부내역 합 490,070원, 합계(B) 490,060원. "
     "또한 본당 490,060 + 자체 60 = 490,120원 ≠ 합계(B) 490,060원", "미해결"),
    ([32], "누계 누락", "신청누계액 1,300,000원 과소(신앙학교 준비비 누락). 계산값 14,400,655원",
     "4_6월교리준비비_추가신청서(미완).hwpx 누계 13,100,655원. 계산: 직전 누계 14,171,060 + 259,300 "
     "− 6월 교리준비비 반납 29,705 = 14,400,655원. 차이 1,300,000원 = (예)2026-031 신앙학교 준비비", "미해결"),
    ([36], "기재 오류", "잔액란 공란",
     "8_6월교리준비비_보고서(추가보고서).hwp 잔액(A-B)란 공란", "미해결"),
    ([36, 37], "번호 중복", "문서번호 036이 6월 교리준비비 추가보고서와 7월 운영비 신청서에 중복",
     "(집)2026-036 = 8_6월교리준비비_보고서(추가보고서).hwp, (예)2026-036 = 1_7월교사회운영비_신청서.hwpx", "미해결"),
    ([37], "계산 불일치", "신청누계액이 계산값(14,880,655원)과 2,850원 차이",
     "1_7월교사회운영비_신청서.hwpx 누계 14,883,505원. 계산: 14,400,655 + 480,000 = 14,880,655원", "미해결"),
    ([37, 38], "계산 불일치", "잔여액이 배정액 − 신청누계액과 480,000원 차이",
     "이관 중 발견. 36,470,000 − 14,883,505 = 21,586,495원이나 잔여액 22,066,495원 기재"
     "(금번 신청액 반영 전 누계로 계산한 값). 부서 운영비 신청서도 같은 값을 그대로 기재. "
     "두 문서 잔여액이 자체 누계 기준 계산값보다 정확히 480,000원 큼 — 이전 문서 값 복사 가능성",
     "미해결", ["I-018"]),
    ([38], "누계 누락", "부서 운영비 480,000원 신청누계 미반영(이후 문서 모두 제외)",
     "2_부서 운영비_신청서.hwpx 누계 14,883,505원(7월 운영비 신청서와 동일). "
     "이후 (예)039~052 누계가 모두 480,000원(−2,850원 오차 포함 477,150원) 적게 기재됨", "미해결"),
    ([38], "문서 누락", "집행보고서 없음, 지급 여부 미확인(1~8월분 60,000원×8개월 일괄 신청 추정)",
     "2_부서 운영비_신청서.hwpx 세부내역 60,000 × 8 = 480,000원. 대응 집행보고서 없음. "
     "사업계획 P13 부서 운영비 월 60,000원", "미해결"),
    ([38], "미확인", "계정과목 '교사회 운영비 60,000×8' — 부서 운영비 월 60,000원 × 8개월(1~8월분) "
     "일괄 신청일 가능성. 기재 오류 여부 확인 필요",
     "2_부서 운영비_신청서.hwpx 계정과목란 '교사회 운영비'(사업은 P13 부서 운영비)", "미해결"),
    ([39], "계산 불일치", "세부내역 합 593,400원인데 합계 480,000원 기재, 자체 103,400원 반영 시에도 10,000원 차이",
     "3_7월교사회운영비_보고서.hwp 세부합 593,400 / 합계(B) 480,000 / 자체 103,400 / 본당 480,000. "
     "본당 + 자체 = 583,400원으로 세부합과 10,000원 차이", "미해결"),
    ([53], "기재 오류", "행사명이 '9월 교사회 운영비'로 기재됨(10월분)",
     "1_10월교사회운영비_신청서.hwpx 행사명 '9월 교사회 운영비', 파일명·작성일(10월 10일)은 10월분", "미해결"),
]


def map_project(name: str):
    for pattern, pid, confident in PROJECT_RULES:
        if re.search(pattern, name):
            return pid, confident
    return "", False


def read_ledger_xlsx(path) -> list[dict]:
    ws = openpyxl.load_workbook(path, data_only=True).worksheets[0]
    rows = ws.iter_rows(values_only=True)
    header = [LEDGER_HEADER_MAP[nfc(h)] for h in next(rows) if h is not None]
    out = []
    for values in rows:
        r = dict(zip(header, values))
        if r.get("구분") not in ("예산신청서", "집행보고서"):
            continue  # 하단 ※ 주석 등
        out.append(r)
    return out


def pair_documents(ledger: list[dict]) -> None:
    """신청서마다 같은 월·사업ID의, 뒤에 오는 첫 미연결 보고서를 짝짓는다."""
    def ref(r):
        no = r["문서번호"]
        return no if no and "미기재" not in no else f"#{r['순번']}"

    used = set()
    for req in (r for r in ledger if r["구분"] == "예산신청서"):
        for rep in ledger:
            if (rep["구분"] == "집행보고서" and rep["순번"] > req["순번"]
                    and rep["순번"] not in used and rep["월"] == req["월"]
                    and rep["사업ID"] and rep["사업ID"] == req["사업ID"]):
                used.add(rep["순번"])
                req["연결문서번호"], rep["연결문서번호"] = ref(rep), ref(req)
                d1, d2 = parse_date(req["작성일"]), parse_date(rep["작성일"])
                state = "사후정산" if d1 and d1 == d2 else "정산완료"
                req["상태"] = rep["상태"] = state
                break

    last_month = ledger[-1]["월"]
    for r in ledger:
        if r["상태"]:
            continue
        if r["구분"] == "예산신청서":
            r["상태"] = "진행중" if r["월"] == last_month else "보고서없음"
        else:
            r["상태"] = "자체집행" if not r["본당보조"] else "신청서없음"


def build_ledger(src) -> tuple[list[dict], list[dict]]:
    raw = read_ledger_xlsx(find_child(src, LEDGER_XLSX))
    ledger, remarks = [], []
    for seq, r in enumerate(raw, 1):
        row = {c: r.get(c) for c in LEDGER_COLUMNS}
        row["순번"] = seq
        row["작성일"] = r.get("작성일") or ""
        row["사업ID"], confident = map_project(nfc(r["행사명"]))
        row["연결문서번호"] = ""
        row["상태"] = ""
        row["_confident"] = confident
        ledger.append(row)
        if r.get("비고"):
            remarks.append((seq, nfc(r["비고"])))
    pair_documents(ledger)
    return ledger, remarks


def build_plan(src) -> tuple[list[dict], list[dict]]:
    wb = openpyxl.load_workbook(find_child(src, PLAN_XLSX), data_only=False)
    ws = wb["사업계획"]
    months = [f"{m}월" for m in range(1, 13)]
    plan = []
    header_row = None
    for row in ws.iter_rows(values_only=True):
        if row[0] == "사업ID":
            header_row = [nfc(h) if h else "" for h in row]
            continue
        if header_row is None or not (isinstance(row[0], str) and re.fullmatch(r"P\d{2}", row[0])):
            continue
        r = dict(zip(header_row, row))
        k = lambda v: int(v or 0) * 1000  # 천원 → 원
        p = {
            "사업ID": r["사업ID"], "시기": r["시기"], "계속신규": r["계속/신규"],
            "사업명": r["사업명"], "세부내역": r["세부내역"],
            "자체예산": k(r["자체"]), "본당예산": k(r["본당"]),
            "전년도집행액": "" if r["전년도 본당 집행액"] is None else k(r["전년도 본당 집행액"]),
        }
        for m in months:
            p[m] = k(r[m])
        plan.append(p)

    info = [{"항목": nfc(a), "내용": nfc(b)}
            for a, b in wb["개요·비고"].iter_rows(min_row=2, values_only=True) if a]
    return plan, info


def build_issues(ledger: list[dict]) -> list[dict]:
    by_seq = {r["순번"]: r for r in ledger}
    issues = []
    for n, (seqs, kind, desc, basis, state, *rest) in enumerate(INITIAL_ISSUES, 1):
        related = rest[0] if rest else []
        issues.append({
            "이슈ID": f"I-{n:03d}",
            "대상순번": MULTI_SEP.join(map(str, seqs)),
            "대상문서번호": MULTI_SEP.join(by_seq[s]["문서번호"] or "" for s in seqs),
            "유형": kind, "설명": desc, "근거": basis,
            "관련이슈": MULTI_SEP.join(related), "상태": state,
        })
    return issues


def build_settings() -> dict:
    return {
        "기본정보": {
            "본당명": "천주교 서울대교구 개포동성당",
            "회계연도": 2026,
            "분과": "청소년분과",
            "주관단체": "초등부 교사회",
            "분과장": "임인호",
            "기본장소": "개포동성당",
            "참가대상": "초등부 교사회",
        },
        "기안정보": {
            "기안자": "정연재 레오",
            "기안단체장": "정지연 플로라",
        },
        "금액문서기준": {
            "예산배정액": 36470000,
            "운영비단가": 60000,
            "교사수": 9,
            "결재란": ["분과장", "재정분과장", "사목회장", "담당사제", "사무장"],
            "신청서번호접두": "(예)2026-",
            "보고서번호접두": "(집)2026-",
            "다음문서번호": 53,
            "기본시작시각": "10:00",
            "기본종료시각": "18:00",
        },
        # 2026년/8월/7_교사하계워크숍비_보고서.hwp 참가자 (신앙캠프 보고서에는 명단 없음)
        "교사명단": [
            "박지우 엘리사벳", "이유진 라파엘라", "이찬주 바오로", "양준형 사도요한",
            "양희강 스테파노", "정연재 레오", "정지연 플로라", "정하영 실비아", "조민기 사도요한",
        ],
        "모델": "claude-sonnet-5-5",
    }


def main(argv=None) -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--force", action="store_true")
    args = ap.parse_args(argv)

    targets = [LEDGER_CSV, PLAN_CSV, PLAN_INFO_CSV, ISSUES_CSV, HISTORY_CSV, SETTINGS_JSON]
    existing = [p.name for p in targets if p.exists()]
    if existing and not args.force:
        print(f"이미 존재: {', '.join(existing)} — 덮어쓰려면 --force", file=sys.stderr)
        return 1

    src = source_dir(2026)
    ledger, remarks = build_ledger(src)
    plan, info = build_plan(src)
    issues = build_issues(ledger)

    DATA_DIR.mkdir(exist_ok=True)
    write_csv(LEDGER_CSV, LEDGER_COLUMNS, ledger)
    write_csv(PLAN_CSV, list(plan[0].keys()), plan)
    write_csv(PLAN_INFO_CSV, ["항목", "내용"], info)
    write_csv(ISSUES_CSV, ISSUE_COLUMNS, issues)
    write_csv(HISTORY_CSV, HISTORY_COLUMNS, [])
    SETTINGS_JSON.write_text(json.dumps(build_settings(), ensure_ascii=False, indent=2) + "\n",
                             encoding="utf-8")

    # ---- 요약 ----
    print(f"장부 {len(ledger)}행 (신청서 {sum(r['구분'] == '예산신청서' for r in ledger)}, "
          f"보고서 {sum(r['구분'] == '집행보고서' for r in ledger)})")
    print("\n[사업ID 매핑]")
    for pid in sorted({r["사업ID"] for r in ledger}):
        seqs = [str(r["순번"]) for r in ledger if r["사업ID"] == pid]
        print(f"  {pid or '(없음)'}: {len(seqs)}건  순번 {','.join(seqs)}")
    print("\n[확인 필요]")
    for r in ledger:
        if not r["_confident"]:
            print(f"  순번 {r['순번']} {r['문서번호']} '{r['행사명']}' → 제안 {r['사업ID'] or '-'}")
    print("\n[짝 없는 문서]")
    for r in ledger:
        if not r["연결문서번호"]:
            print(f"  순번 {r['순번']} {r['구분']} {r['문서번호']} '{r['행사명']}' 상태={r['상태']}")
    print("\n[사후정산 쌍]")
    for r in ledger:
        if r["상태"] == "사후정산" and r["구분"] == "예산신청서":
            print(f"  {r['문서번호']} ↔ {r['연결문서번호']} {r['작성일']} '{r['행사명']}'")
    print(f"\n[이슈] {len(issues)}건")
    for t in sorted({i['유형'] for i in issues}):
        print(f"  {t}: {sum(i['유형'] == t for i in issues)}")
    targeted = {int(s) for i in issues for s in i["대상순번"].split(MULTI_SEP)}
    missing = [seq for seq, _ in remarks if seq not in targeted]
    print(f"\n엑셀 비고 {len(remarks)}행 → 이슈 미등록 {missing or '없음'}")
    return 1 if missing else 0


if __name__ == "__main__":
    sys.exit(main())
