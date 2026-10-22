"""영수증 첨부 문서(docx). 제목(문서번호 + 행사명) 아래 순번별로 [이미지 → 날짜·상호·금액·용도].

이미지는 jpg/png/heic 또는 pdf(첫 쪽)를 받는다. EXIF 방향을 바로잡고, 흰 여백을 잘라 내고,
페이지 폭(그리고 본문 높이의 일정 비율)에 맞춰 줄인다.
"""
import io
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Emu, Pt
from PIL import Image, ImageChops, ImageOps

from core.models import Receipt

FONT = "맑은 고딕"
MAX_PIXELS = 1600            # 긴 변 최대 픽셀(파일 크기 제한)
IMAGE_HEIGHT_RATIO = 0.62    # 본문 높이 중 이미지가 차지할 최대 비율(설명 표와 한 쪽에)
PDF_SCALE = 2.0

try:  # 아이폰 HEIC
    from pillow_heif import register_heif_opener
    register_heif_opener()
except ImportError:  # pragma: no cover
    pass


def load_image(path) -> Image.Image:
    path = Path(path)
    if path.suffix.lower() == ".pdf":
        import pypdfium2 as pdfium
        pdf = pdfium.PdfDocument(str(path))
        try:
            return pdf[0].render(scale=PDF_SCALE).to_pil().convert("RGB")
        finally:
            pdf.close()
    with Image.open(path) as im:
        return ImageOps.exif_transpose(im).convert("RGB")


def trim_margin(im: Image.Image, pad: int = 12, threshold: int = 12) -> Image.Image:
    """가장자리와 같은 색(대개 흰색)의 여백을 잘라 낸다."""
    bg = Image.new("RGB", im.size, im.getpixel((0, 0)))
    diff = ImageChops.difference(im, bg).convert("L").point(lambda v: 255 if v > threshold else 0)
    box = diff.getbbox()
    if not box:
        return im
    l, t, r, b = box
    return im.crop((max(0, l - pad), max(0, t - pad), min(im.width, r + pad), min(im.height, b + pad)))


def prepare_image(path) -> Image.Image:
    im = trim_margin(load_image(path))
    im.thumbnail((MAX_PIXELS, MAX_PIXELS))
    return im


def _font(run, size=None, bold=False, font=FONT):
    run.font.name = font
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), font)
    if size:
        run.font.size = Pt(size)
    run.bold = bold


def _fit(im: Image.Image, max_w: int, max_h: int) -> dict:
    ratio = im.width / im.height
    if max_w / ratio <= max_h:
        return {"width": Emu(max_w)}
    return {"height": Emu(max_h)}


def build_receipts_docx(문서번호: str, 행사명: str, receipts: list[Receipt], path,
                        base_dir=None, font: str = FONT) -> Path:
    """font: 한글 글꼴. 기본 맑은 고딕(윈도우·한글 환경)."""
    def text(run, size=None, bold=False):
        _font(run, size, bold, font)

    doc = Document()
    sec = doc.sections[0]
    body_w = sec.page_width - sec.left_margin - sec.right_margin
    body_h = sec.page_height - sec.top_margin - sec.bottom_margin

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    text(title.add_run(f"{문서번호}  {행사명}"), 15, bold=True)
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    total = sum(r.금액 or 0 for r in receipts)
    text(sub.add_run(f"영수증 첨부 {len(receipts)}건 · 합계 {total:,}원"), 10)

    for n, rc in enumerate(sorted(receipts, key=lambda r: r.순번)):
        src = Path(base_dir) / rc.이미지 if base_dir else Path(rc.이미지)
        im = prepare_image(src)
        buf = io.BytesIO()
        im.save(buf, format="JPEG", quality=85)
        buf.seek(0)

        head = doc.add_paragraph()
        head.paragraph_format.keep_with_next = True
        if n:
            head.paragraph_format.page_break_before = True
        text(head.add_run(f"{rc.순번}."), 12, bold=True)

        pic = doc.add_paragraph()
        pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pic.paragraph_format.keep_with_next = True
        pic.add_run().add_picture(buf, **_fit(im, int(body_w), int(body_h * IMAGE_HEIGHT_RATIO)))

        table = doc.add_table(rows=4, cols=2)
        table.style = "Table Grid"
        table.autofit = False
        widths = (Emu(int(body_w * 0.18)), Emu(int(body_w * 0.82)))
        for col, w in zip(table.columns, widths):
            col.width = w
        rows = (("날짜", rc.날짜.isoformat() if rc.날짜 else ""), ("상호", rc.상호),
                ("금액", f"{rc.금액:,}원" if rc.금액 is not None else ""), ("용도", rc.용도))
        for row, (label, value) in zip(table.rows, rows):
            for cell, w in zip(row.cells, widths):
                cell.width = w
            text(row.cells[0].paragraphs[0].add_run(label), 10, bold=True)
            text(row.cells[1].paragraphs[0].add_run(value), 10)

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    return path
